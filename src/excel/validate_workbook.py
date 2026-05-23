"""Workbook structural validation helpers.

These checks do not replace native Excel calculation. They catch the model errors
that kept recurring during generated workbook creation: stale references,
#REF-style formula text, invalid scenario selectors, missing required sheets and
formula links pointing to the wrong modelling rows.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

REQUIRED_SHEETS = [
    "Assumptions",
    "3 Statement Model",
    "WACC",
    "DCF",
    "Comps Analysis",
    "Football Field",
    "Sensitivity Analysis",
]

FORMULA_ERROR_TOKENS = ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]
STALE_REFERENCE_TOKENS = [
    "B$49",
    "$B$65",
    "B65",
    "B71-'Assumptions'!B65",
    "'3 Statement Model'!J44",  # old DCF link to ending cash instead of FCF
]


@dataclass
class WorkbookValidationResult:
    ok: bool
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _formula_cells(workbook_path: str | Path):
    wb = load_workbook(workbook_path, data_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    yield ws.title, cell.coordinate, cell.value


def validate_workbook_structure(workbook_path: str | Path) -> WorkbookValidationResult:
    path = Path(workbook_path)
    result = WorkbookValidationResult(ok=True)
    wb = load_workbook(path, data_only=False)

    for sheet in REQUIRED_SHEETS:
        if sheet not in wb.sheetnames:
            result.errors.append(f"Missing required sheet: {sheet}")

    if result.errors:
        result.ok = False
        return result

    assumptions = wb["Assumptions"]
    if assumptions["B3"].value not in {"Bear", "Base", "Bull"}:
        result.errors.append("Assumptions!B3 must default to one of Bear, Base or Bull.")

    dcf = wb["DCF"]
    if dcf["J14"].value != "=J12/J13":
        result.errors.append("DCF!J14 should be the target price formula =J12/J13.")
    if not str(dcf["B5"].value).startswith("='3 Statement Model'!") or not str(dcf["B5"].value).endswith("45"):
        result.errors.append("DCF!B5 should link to Free Cash Flow row 45, not Ending Cash row 44.")

    ff = wb["Football Field"]
    if ff["C5"].value != "=DCF!J14":
        result.errors.append("Football Field!C5 should link to the DCF target price.")

    for sheet, cell, formula in _formula_cells(path):
        for token in FORMULA_ERROR_TOKENS:
            if token in formula:
                result.errors.append(f"Formula error token {token} in {sheet}!{cell}: {formula}")
        for token in STALE_REFERENCE_TOKENS:
            if token in formula:
                result.errors.append(f"Stale reference token {token} in {sheet}!{cell}: {formula}")

    result.ok = not result.errors
    return result


def raise_if_invalid(workbook_path: str | Path) -> None:
    result = validate_workbook_structure(workbook_path)
    if not result.ok:
        details = "\n".join(result.errors[:20])
        raise ValueError(f"Generated workbook failed validation:\n{details}")
