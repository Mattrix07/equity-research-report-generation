"""Evidence overlay for generated Excel models.

Adds source-driven evidence to the generated workbook after the main model is
created. This keeps the modelling engine separate from the evidence extraction
layer and allows the LLM/source-agent layer to populate sector-specific tabs.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from src.schemas import FullReport

DARK_BLUE = "17365D"
WHITE = "FFFFFF"
GREY = "D9E1F2"
YELLOW = "FFF2CC"
BLUE = "0000FF"


def _section(ws, row: int, title: str, end_col: int = 10) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    c = ws.cell(row, 1, title)
    c.fill = PatternFill("solid", fgColor=DARK_BLUE)
    c.font = Font(color=WHITE, bold=True)
    return row + 1


def _headers(ws, row: int, labels: list[str]) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row, i, label)
        c.fill = PatternFill("solid", fgColor=GREY)
        c.font = Font(bold=True)
        c.border = Border(bottom=Side(style="thin", color="808080"))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    widths = {"A": 20, "B": 28, "C": 28, "D": 24, "E": 20, "F": 18, "G": 50, "H": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _write_rows(ws, start_row: int, rows: list[dict[str, Any]], preferred_columns: list[str] | None = None) -> int:
    if not rows:
        ws.cell(start_row, 1, "No company-specific evidence available yet.")
        return start_row + 2

    columns = preferred_columns or []
    discovered = []
    for row in rows:
        for key in row.keys():
            if key not in discovered:
                discovered.append(key)
    columns = columns + [c for c in discovered if c not in columns]
    columns = columns[:10]

    _headers(ws, start_row, [c.replace("_", " ").title() for c in columns])
    for r_idx, row in enumerate(rows, start=start_row + 1):
        for c_idx, col in enumerate(columns, start=1):
            cell = ws.cell(r_idx, c_idx, row.get(col, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return start_row + len(rows) + 3


def _create_evidence_summary(wb, report: FullReport) -> None:
    if "Evidence Summary" in wb.sheetnames:
        del wb["Evidence Summary"]
    ws = wb.create_sheet("Evidence Summary", 1)
    _style(ws)
    ws["A1"] = f"Evidence Summary - {report.snapshot.company_name} ({report.snapshot.ticker})"
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    ws.merge_cells("A1:H1")

    pack = report.company_evidence
    r = 3
    r = _section(ws, r, "Evidence Items", 8)
    _headers(ws, r, ["Source Type", "Source", "Title", "Topic", "Value", "Period", "URL", "Excerpt"])
    if pack and pack.evidence_items:
        for i, item in enumerate(pack.evidence_items, start=r + 1):
            values = [
                item.source_type,
                item.source_name,
                item.title,
                item.metric_or_topic,
                item.value,
                item.period_or_date,
                item.url,
                item.excerpt,
            ]
            for j, value in enumerate(values, start=1):
                cell = ws.cell(i, j, value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if j in {5, 7}:
                    cell.font = Font(color=BLUE)
        r = r + len(pack.evidence_items) + 3
    else:
        ws.cell(r + 1, 1, "No evidence pack generated.")
        r += 4

    r = _section(ws, r, "Evidence Gaps and Warnings", 8)
    _headers(ws, r, ["Type", "Item"])
    row = r + 1
    if pack:
        for gap in pack.gaps[:30]:
            ws.cell(row, 1, "Gap")
            ws.cell(row, 2, gap)
            row += 1
        for warning in pack.warnings[:30]:
            ws.cell(row, 1, "Warning")
            ws.cell(row, 2, warning)
            row += 1
    else:
        ws.cell(row, 1, "Gap")
        ws.cell(row, 2, "No source evidence pack found.")


def _append_sector_evidence(wb, report: FullReport) -> None:
    pack = report.company_evidence
    if not pack or not pack.sector_tab_evidence:
        return

    for tab_name, tab_evidence in pack.sector_tab_evidence.items():
        if tab_name not in wb.sheetnames:
            ws = wb.create_sheet(tab_name[:31])
        else:
            ws = wb[tab_name]
        for col in range(1, 12):
            ws.column_dimensions[get_column_letter(col)].width = max(ws.column_dimensions[get_column_letter(col)].width or 12, 18)

        start_row = max(ws.max_row + 3, 20)
        start_row = _section(ws, start_row, "Company-Specific Evidence Overlay", 10)
        if tab_evidence.notes:
            for note in tab_evidence.notes:
                ws.cell(start_row, 1, "Note")
                ws.cell(start_row, 2, note)
                ws.cell(start_row, 2).alignment = Alignment(wrap_text=True)
                start_row += 1
            start_row += 1
        _write_rows(ws, start_row, tab_evidence.rows)


def apply_evidence_overlay(report: FullReport, workbook_path: str) -> str:
    path = Path(workbook_path)
    wb = load_workbook(path)
    _create_evidence_summary(wb, report)
    _append_sector_evidence(wb, report)
    wb.save(path)
    return str(path)
