"""Template-controlled Excel exporter.

This is the production export path for the simplified model. It keeps the model
architecture fixed, validates the workbook structure, and optionally uses native
Excel via xlwings to recalculate before the file is returned.

Why this exists:
- LLMs should not dynamically invent Excel formulas.
- The workbook layout should be stable and template-like.
- Python should populate inputs and validate links.
- Excel, when available, should perform the actual formula calculation.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.excel.excel_runtime import recalc_with_excel_if_enabled
from src.excel.validate_workbook import raise_if_invalid
from src.report.repaired_simplified_excel_exporter import export_excel_model as export_repaired_model
from src.schemas import FullReport


def _write_runtime_status(workbook_path: str | Path, status: str) -> None:
    wb = load_workbook(workbook_path)
    ws = wb["Assumptions"]
    ws["D3"] = "Excel Runtime Status"
    ws["E3"] = status
    ws["E3"].comment = None
    wb.save(workbook_path)


def export_excel_model(report: FullReport, output_dir: Path) -> str:
    """Create, validate, optionally recalculate and return the model path."""
    path = export_repaired_model(report, output_dir)

    # First structural validation before Excel runtime. This catches broken
    # formulas before a user downloads the file.
    raise_if_invalid(path)

    recalculated, runtime_message = recalc_with_excel_if_enabled(path)
    _write_runtime_status(path, runtime_message)

    # Re-run validation after runtime status is written and, where available,
    # after native Excel has recalculated the workbook.
    raise_if_invalid(path)
    return str(path)
